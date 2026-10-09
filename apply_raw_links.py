from docx import Document

links = [
    "https://scholar.google.com/scholar?q=Perancangan+dan+Analisis+Mekanika+Permainan+2D+Arcade+Brick+Breaker",
    "https://www.construct.net/en/make-games/manuals/construct-3",
    "https://scholar.google.com/scholar?q=Pemanfaatan+Platform+Desain+Canva+dalam+Perancangan+Aset+Visual+User+Interface",
    "https://scholar.google.com/scholar?q=Pemodelan+Sistem+Informasi+dan+Alur+Kerja+Perangkat+Lunak+Berbasis+Unified+Modelling+Language+(UML)",
    "https://scholar.google.com/scholar?q=Logika+Algoritma+Pemrograman+dan+Dokumentasi+Diagram+Alir+(Flowchart)",
    "https://scholar.google.com/scholar?q=Penerapan+Data+Flow+Diagram+(DFD)+dalam+Analisis+Aliran+Data",
    "https://scholar.google.com/scholar?q=Analisis+Kebutuhan+Fungsional+Pengguna+Berbasis+Use+Case+Diagram"
]

def process_doc(input_file, output_file):
    try:
        doc = Document(input_file)
        found = False
        idx = 0
        for p in doc.paragraphs:
            if "DAFTAR PUSTAKA" in p.text and p.style.name.startswith("Heading"):
                found = True
                continue
            
            if found:
                if p.style.name.startswith("Heading") and p.text.strip():
                    break
                text = p.text.strip()
                if text and text.startswith("["):
                    # 1. Clean up old hyperlink and text
                    for child in list(p._p):
                        if child.tag.endswith('hyperlink'):
                            p._p.remove(child)
                            
                    for r in p.runs:
                        if "(Tautan: " in r.text:
                            r.text = r.text.replace(" (Tautan: ", "")
                        if ")" in r.text and len(r.text.strip()) == 1:
                            r.text = r.text.replace(")", "")
                            
                    # 2. Add raw text link
                    if idx < len(links):
                        p.add_run(f" Tersedia di: {links[idx]}")
                        idx += 1
                        
        doc.save(output_file)
        print(f"Success for {output_file}")
    except Exception as e:
        print(f"Error for {output_file}: {e}")

process_doc("LAPORAN_PW_FAIRUZ_REVISI_20 Mei (2) (1)_REVISED.docx", "LAPORAN_PW_FAIRUZ_REVISI_20 Mei (2) (1)_REVISED.docx")
process_doc("LAPORAN_PW_FAIRUZ_TEMP_v3.docx", "LAPORAN_PW_FAIRUZ_TEMP_v4.docx")
