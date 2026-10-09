import docx
from docx import Document
from docx.oxml.shared import OxmlElement, qn
from docx.opc.constants import RELATIONSHIP_TYPE

def add_hyperlink(paragraph, text, url):
    # Get the paragraph part
    part = paragraph.part
    # Create the relationship for the hyperlink
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    # Create the <w:hyperlink> element
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    # Create the run element <w:r>
    new_run = OxmlElement('w:r')
    
    # Create the run properties <w:rPr>
    rPr = OxmlElement('w:rPr')
    
    # Create the color element <w:color w:val="0563C1"/>
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0563C1')
    rPr.append(c)

    # Create underline element <w:u w:val="single"/>
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)

    new_run.append(rPr)

    # Create text element <w:t>
    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)
    
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

doc = Document("LAPORAN_PW_FAIRUZ_TEMP_v2.docx")

links = [
    "https://scholar.google.com/scholar?q=Perancangan+dan+Analisis+Mekanika+Permainan+2D+Arcade+Brick+Breaker",
    "https://www.construct.net/en/make-games/manuals/construct-3",
    "https://scholar.google.com/scholar?q=Pemanfaatan+Platform+Desain+Canva+dalam+Perancangan+Aset+Visual+User+Interface",
    "https://scholar.google.com/scholar?q=Pemodelan+Sistem+Informasi+dan+Alur+Kerja+Perangkat+Lunak+Berbasis+Unified+Modelling+Language+(UML)",
    "https://scholar.google.com/scholar?q=Logika+Algoritma+Pemrograman+dan+Dokumentasi+Diagram+Alir+(Flowchart)",
    "https://scholar.google.com/scholar?q=Penerapan+Data+Flow+Diagram+(DFD)+dalam+Analisis+Aliran+Data",
    "https://scholar.google.com/scholar?q=Analisis+Kebutuhan+Fungsional+Pengguna+Berbasis+Use+Case+Diagram"
]

found = False
idx = 0
for p in doc.paragraphs:
    if "DAFTAR PUSTAKA" in p.text and p.style.name.startswith("Heading"):
        found = True
        continue
    
    if found:
        if p.style.name.startswith("Heading") and p.text.strip():
            break
            
        text = p.text.strip()
        if text and text.startswith("["):
            if idx < len(links):
                # Add a space and the link
                p.add_run(" (URL: ")
                add_hyperlink(p, "Klik di sini", links[idx])
                p.add_run(")")
                idx += 1

doc.save("LAPORAN_PW_FAIRUZ_TEMP_v2.docx")
doc.save("LAPORAN_PW_FAIRUZ_REVISI_20 Mei (2) (1)_REVISED.docx")
print("Done adding links")
