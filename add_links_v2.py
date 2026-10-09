import docx
from docx import Document
from docx.oxml.shared import OxmlElement, qn
from docx.opc.constants import RELATIONSHIP_TYPE

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0563C1')
    rPr.append(c)

    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)

    new_run.append(rPr)

    text_elem = OxmlElement('w:t')
    text_elem.text = text
    new_run.append(text_elem)
    
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

links = [
    "https://scholar.google.com/scholar?q=Perancangan+dan+Analisis+Mekanika+Permainan+2D+Arcade+Brick+Breaker",
    "https://www.construct.net/en/make-games/manuals/construct-3",
    "https://scholar.google.com/scholar?q=Pemanfaatan+Platform+Desain+Canva+dalam+Perancangan+Aset+Visual+User+Interface",
    "https://scholar.google.com/scholar?q=Pemodelan+Sistem+Informasi+dan+Alur+Kerja+Perangkat+Lunak+Berbasis+Unified+Modelling+Language+(UML)",
    "https://scholar.google.com/scholar?q=Logika+Algoritma+Pemrograman+dan+Dokumentasi+Diagram+Alir+(Flowchart)",
    "https://scholar.google.com/scholar?q=Penerapan+Data+Flow+Diagram+(DFD)+dalam+Analisis+Aliran+Data",
    "https://scholar.google.com/scholar?q=Analisis+Kebutuhan+Fungsional+Pengguna+Berbasis+Use+Case+Diagram"
]

def process_doc(doc_path, save_path):
    try:
        doc = Document(doc_path)
        found = False
        idx = 0
        for p in doc.paragraphs:
            if "DAFTAR PUSTAKA" in p.text and p.style.name.startswith("Heading"):
                found = True
                continue
            
            if found:
                if p.style.name.startswith("Heading") and p.text.strip():
                    break
                    
                text = p.text.strip()
                if text and text.startswith("["):
                    if idx < len(links):
                        p.add_run(" (Tautan: ")
                        add_hyperlink(p, "Buka Sumber", links[idx])
                        p.add_run(")")
                        idx += 1
        
        doc.save(save_path)
        print(f"Successfully processed {save_path}")
    except Exception as e:
        print(f"Error processing {doc_path} -> {save_path}: {e}")

process_doc("LAPORAN_PW_FAIRUZ_TEMP_v2.docx", "LAPORAN_PW_FAIRUZ_TEMP_v3.docx")
process_doc("LAPORAN_PW_FAIRUZ_REVISI_20 Mei (2) (1)_REVISED.docx", "LAPORAN_PW_FAIRUZ_REVISI_20 Mei (2) (1)_REVISED.docx")
